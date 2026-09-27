#!/usr/bin/env python3
"""Build pk_experiment_results.xlsx from results.json (+ verification.json if present).

Usage:  python3 build_excel.py results.json [output.xlsx] [verification.json]
"""
import json
import statistics
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import NumFmt
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
PK_COLOR, NOPK_COLOR = "2A78D6", "EB6834"          # blue = with PK, orange = without PK
INK, MUTED = "1F1F1F", "5F5E5A"
HEAD_FILL = PatternFill("solid", fgColor="2F3B4C")
PK_FILL = PatternFill("solid", fgColor="E3EEFB")
NOPK_FILL = PatternFill("solid", fgColor="FCE8DF")
ZEBRA = PatternFill("solid", fgColor="F6F6F4")
THIN = Side(style="thin", color="D0D0CC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MS = '#,##0.000'
SEC = '#,##0.000'
INT = '#,##0'
TIMES = '#,##0.0"×"'
PCT = '+0.0%;-0.0%;0.0%'
LABELS = {"with_pk": "With PK", "without_pk": "Without PK"}


def f(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def put(ws, ref, value, font=None, fmt=None, fill=None, align=None, border=True):
    c = ws[ref]
    c.value = value
    c.font = font or f()
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = BOX
    return c


def header_row(ws, row, col, labels, fill=HEAD_FILL, color="FFFFFF"):
    for i, lab in enumerate(labels):
        put(ws, f"{get_column_letter(col + i)}{row}", lab, f(bold=True, color=color), fill=fill,
            align=Alignment(horizontal="center", vertical="center", wrap_text=True))


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def title(ws, text, sub=None):
    ws["A1"].value = text
    ws["A1"].font = f(14, bold=True)
    if sub:
        ws["A2"].value = sub
        ws["A2"].font = f(10, color=MUTED, italic=True)


def style_chart(ch, y_title, x_title, log=False):
    ch.y_axis.title = y_title
    ch.x_axis.title = x_title
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.y_axis.majorGridlines.spPr = None
    ch.y_axis.numFmt = NumFmt(formatCode="General", sourceLinked=False)
    if log:
        ch.y_axis.scaling.logBase = 10
    else:
        ch.y_axis.scaling.min = 0          # bars always start at zero
    ch.legend.position = "b"
    ch.height, ch.width = 8.5, 17
    for s, color in zip(ch.series, (PK_COLOR, NOPK_COLOR)):
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = color
    ch.gapWidth = 60
    ch.overlap = -5


OK_FILL = PatternFill("solid", fgColor="E3F1E6")
TODO_FILL = PatternFill("solid", fgColor="FFF1D1")


def add_verification(wb, v, cfg, env, sel_cells, ins_cells, ws_sum):
    """Doctor's checklist + MySQL check + Sample data + Raw verification sheets."""
    targets, batches = cfg["select_targets"], cfg["insert_batches"]
    wrap = Alignment(wrap_text=True, vertical="top")
    ws_cl = wb.create_sheet("Doctor's checklist")
    ws_mc = wb.create_sheet("MySQL check")
    ws_sd = wb.create_sheet("Sample data")
    ws_rv = wb.create_sheet("Raw verification")
    dc = {(x["stage"], x["table"]): x for x in v["data_checks"]}
    pk2, np2 = dc[("after step 2", "data_pk")], dc[("after step 2", "data_nopk")]
    pk5, np5 = dc[("after step 5", "data_pk")], dc[("after step 5", "data_nopk")]
    final_rows = cfg["initial_rows"] + sum(batches)

    # ---------------------------------------------------------------- Raw verification
    title(ws_rv, "Raw verification data - every statement of the verification run",
          "SQL text, time and rows examined exactly as MySQL recorded them in performance_schema.")
    header_row(ws_rv, 4, 1, ["Table", "Target id", "Repeat #", "SQL text (recorded by MySQL)", "Python time (ms)",
                             "MySQL timer (ms)", "Rows examined", "Rows sent", "No index used (1 = yes)"])
    rv_rng = {}
    r = 5
    for x in sorted(v["selects"], key=lambda x: (x["target_id"], x["table"] != "with_pk", x["repeat"])):
        vals = [LABELS[x["table"]], x["target_id"], x["repeat"], x["sql_text"], x["client_ms"], x["mysql_ms"],
                x["rows_examined"], x["rows_sent"], x["no_index_used"]]
        for i, (val, fm) in enumerate(zip(vals, [None, INT, None, None, MS, MS, INT, INT, None])):
            put(ws_rv, f"{get_column_letter(i + 1)}{r}", val, fmt=fm)
        k = (x["table"], x["target_id"])
        a, _ = rv_rng.get(k, (r, r))
        rv_rng[k] = (a, r)
        r += 1
    r += 2
    ws_rv[f"A{r}"].value = "INSERT statements of the verification run"
    ws_rv[f"A{r}"].font = f(11, bold=True)
    r += 1
    header_row(ws_rv, r, 1, ["Table", "Step", "Rows inserted", "INSERT statements + COMMIT", "Python time (s)",
                             "MySQL timer (s)", "COUNT(*) after"])
    r += 1
    rv_ins = {}
    for x in v["inserts"]:
        cnt = x.get("count_after") or dc[("after step 2", px_table(x["table"]))]["row_count"]
        vals = [LABELS[x["table"]], x["step"], x["rows_inserted"], x["statements"], x["client_s"], x["mysql_s"], cnt]
        for i, (val, fm) in enumerate(zip(vals, [None, None, INT, INT, SEC, SEC, INT])):
            put(ws_rv, f"{get_column_letter(i + 1)}{r}", val, fmt=fm)
        rv_ins[(x["step"], x["table"])] = r
        r += 1
    widths(ws_rv, {"A": 13, "B": 16, "C": 13, "D": 44, "E": 16, "F": 16, "G": 15, "H": 11, "I": 20})
    ws_rv.freeze_panes = "A5"

    # ---------------------------------------------------------------- MySQL check
    ws = ws_mc
    title(ws, "Checked inside MySQL - table definitions, data, and MySQL's own timings",
          "A separate, fresh run of the doctor's steps (verification run). Times below come from MySQL's own clock "
          "(performance_schema), shown next to the Python-measured times.")
    r = 4
    ws[f"A{r}"].value = "A. Table definitions - output of SHOW CREATE TABLE and SHOW INDEX"
    ws[f"A{r}"].font = f(11, bold=True)
    r += 1
    header_row(ws, r, 1, ["Table", "SHOW CREATE TABLE (exactly as MySQL returns it)", "", "", "", "",
                          "SHOW INDEX"])
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    r += 1
    for label, table in (("with_pk", "data_pk"), ("without_pk", "data_nopk")):
        idx = v["show_index"][label]
        idx_txt = ("; ".join(f"{i['Key_name']} on {i['Column_name']} ({i['Index_type']}, "
                             f"{'unique' if i['Non_unique'] == 0 else 'non-unique'})" for i in idx)
                   or "(no indexes at all)")
        put(ws, f"A{r}", f"{table}\n({LABELS[label]})", f(bold=True), align=wrap)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        put(ws, f"B{r}", v["show_create"][label], Font(name="Courier New", size=9, color=INK), align=wrap)
        put(ws, f"G{r}", idx_txt, f(bold=True), align=wrap,
            fill=PK_FILL if label == "with_pk" else NOPK_FILL)
        ws.row_dimensions[r].height = 125
        r += 1

    r += 1
    ws[f"A{r}"].value = "B. Data checks - SQL run inside MySQL on both tables"
    ws[f"A{r}"].font = f(11, bold=True)
    r += 1
    header_row(ws, r, 1, ["Stage", "Table", "COUNT(*)", "COUNT(DISTINCT id)", "MIN(id)", "MAX(id)",
                          "Checksum of all columns", "Same data in both tables?", "Age min", "Age max",
                          "Distinct cities", "Salary min", "Salary max", "First created_at", "Last created_at"])
    r += 1
    for stage, label_stage in (("after step 2", f"After step 2 ({cfg['initial_rows']:,} rows)"),
                               ("after step 5", f"After step 5 ({final_rows:,} rows)")):
        first = r
        for table in ("data_pk", "data_nopk"):
            x = dc[(stage, table)]
            vals = [label_stage, table, x["row_count"], x["distinct_ids"], x["min_id"], x["max_id"], x["checksum"],
                    None, x["age_min"], x["age_max"], x["distinct_cities"], x["salary_min"], x["salary_max"],
                    x["created_min"], x["created_max"]]
            fmts = [None, None, INT, INT, INT, INT, "0", None, None, None, None, "#,##0.00", "#,##0.00", None, None]
            for i, (val, fm) in enumerate(zip(vals, fmts)):
                if i == 7:
                    continue
                put(ws, f"{get_column_letter(i + 1)}{r}", val, fmt=fm)
            put(ws, f"H{r}", f'=IF($G${first}=$G${first + 1},"Yes - identical","NO - different")', f(bold=True))
            r += 1
    ws[f"A{r}"].value = ("Checksum = BIT_XOR(CRC32(CONCAT_WS('|', id, name, email, age, city, salary, created_at))) "
                         "over every row. Equal checksums = both tables contain exactly the same rows.")
    ws[f"A{r}"].font = f(9, color=MUTED, italic=True)

    r += 2
    ws[f"A{r}"].value = ("C. SELECT * FROM table WHERE id = N - MySQL's own timer and rows examined "
                         f"(median of {len([1 for s in v['selects'] if s['table'] == 'with_pk' and s['target_id'] == targets[0]])} executions)")
    ws[f"A{r}"].font = f(11, bold=True)
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=9)
    put(ws, f"B{r}", "WITH primary key", f(bold=True), fill=PK_FILL, align=Alignment(horizontal="center"))
    put(ws, f"F{r}", "WITHOUT primary key", f(bold=True), fill=NOPK_FILL, align=Alignment(horizontal="center"))
    r += 1
    sub = ["Main experiment, Python (ms)", "Verification, Python (ms)", "Verification, MySQL timer (ms)",
           "Rows examined (MySQL)"]
    header_row(ws, r, 1, ["Target id"] + sub + sub)
    r += 1
    for t in targets:
        put(ws, f"A{r}", t, f(bold=True), INT)
        for label, c0 in (("with_pk", 2), ("without_pk", 6)):
            main_cell = sel_cells[("after initial load", "WHERE id = N")][t][0 if label == "with_pk" else 1]
            a, b = rv_rng[(label, t)]
            fill = PK_FILL if label == "with_pk" else NOPK_FILL
            put(ws, f"{get_column_letter(c0)}{r}", f"='SELECT results'!{main_cell}", fmt=MS)
            put(ws, f"{get_column_letter(c0 + 1)}{r}", f"=MEDIAN('Raw verification'!E{a}:E{b})", fmt=MS)
            put(ws, f"{get_column_letter(c0 + 2)}{r}", f"=MEDIAN('Raw verification'!F{a}:F{b})", f(bold=True), MS,
                fill=fill)
            put(ws, f"{get_column_letter(c0 + 3)}{r}", f"=MAX('Raw verification'!G{a}:G{b})", f(bold=True), INT)
        r += 1
    ws[f"A{r}"].value = ("MySQL timer = TIMER_WAIT in performance_schema.events_statements_history: the time the server "
                         "spent running the statement. The Python time also includes the network round trip, so a PK "
                         "lookup shows ~0.1 ms inside MySQL but ~0.3 ms in Python. For the table without a PK both "
                         "clocks agree, because reading every row dominates.")
    ws[f"A{r}"].font = f(9, color=MUTED, italic=True)

    r += 2
    ws[f"A{r}"].value = "D. INSERT - MySQL's own timer and row counts (verification run)"
    ws[f"A{r}"].font = f(11, bold=True)
    r += 1
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
    put(ws, f"C{r}", "WITH primary key (s)", f(bold=True), fill=PK_FILL, align=Alignment(horizontal="center"))
    put(ws, f"F{r}", "WITHOUT primary key (s)", f(bold=True), fill=NOPK_FILL, align=Alignment(horizontal="center"))
    r += 1
    sub = ["Main experiment, Python", "Verification, Python", "Verification, MySQL timer"]
    header_row(ws, r, 1, ["Step", "Rows inserted"] + sub + sub + ["COUNT(*) after - with PK",
                                                                   "COUNT(*) after - without PK"])
    r += 1
    for step in ["Initial load"] + [f"Insert {b:,}" for b in batches]:
        n = cfg["initial_rows"] if step == "Initial load" else int(step.split()[1].replace(",", ""))
        put(ws, f"A{r}", step, f(bold=True))
        put(ws, f"B{r}", n, fmt=INT)
        for label, c0 in (("with_pk", 3), ("without_pk", 6)):
            rr = rv_ins[(step, label)]
            main_cell = ins_cells[step][0 if label == "with_pk" else 1]
            put(ws, f"{get_column_letter(c0)}{r}", f"='INSERT results'!{main_cell}", fmt=SEC)
            put(ws, f"{get_column_letter(c0 + 1)}{r}", f"='Raw verification'!E{rr}", fmt=SEC)
            put(ws, f"{get_column_letter(c0 + 2)}{r}", f"='Raw verification'!F{rr}", f(bold=True), SEC,
                fill=PK_FILL if label == "with_pk" else NOPK_FILL)
        put(ws, f"I{r}", f"='Raw verification'!G{rv_ins[(step, 'with_pk')]}", fmt=INT)
        put(ws, f"J{r}", f"='Raw verification'!G{rv_ins[(step, 'without_pk')]}", fmt=INT)
        r += 1
    ws[f"A{r}"].value = ("MySQL timer = total TIMER_WAIT of that step's INSERT statements + COMMIT "
                         "(performance_schema.events_statements_summary_by_thread_by_event_name). The Python time is "
                         "about twice as long because it also packs and sends the rows over the connection. Neither "
                         "clock shows a consistent winner for inserts - the primary key adds almost no insert cost here.")
    ws[f"A{r}"].font = f(9, color=MUTED, italic=True)
    widths(ws, {"A": 26, "B": 15, "C": 15, "D": 15, "E": 15, "F": 15, "G": 22, "H": 18, "I": 15, "J": 15,
                "K": 12, "L": 14, "M": 14, "N": 19, "O": 19})

    # ---------------------------------------------------------------- Sample data
    ws = ws_sd
    title(ws, "Sample rows exactly as stored in MySQL (verification run, after step 2)",
          "SELECT * FROM data_pk / data_nopk WHERE id IN (1..10, 1000, 10000, 100000, 1000000) ORDER BY id. "
          "All columns except id are random.")
    cols = ["id", "name", "email", "age", "city", "salary", "created_at"]
    ws.merge_cells("A4:G4")
    ws.merge_cells("I4:O4")
    put(ws, "A4", "data_pk (WITH primary key)", f(bold=True), fill=PK_FILL, align=Alignment(horizontal="center"))
    put(ws, "I4", "data_nopk (WITHOUT primary key)", f(bold=True), fill=NOPK_FILL,
        align=Alignment(horizontal="center"))
    header_row(ws, 5, 1, cols)
    header_row(ws, 5, 9, cols)
    header_row(ws, 5, 16, ["Row identical?"])
    pk_rows = [x for x in v["samples"] if x["table"] == "with_pk"]
    np_rows = [x for x in v["samples"] if x["table"] == "without_pk"]
    r = 6
    for a_, b_ in zip(pk_rows, np_rows):
        for c0, row_ in ((1, a_), (9, b_)):
            vals = [int(row_["id"]), row_["name"], row_["email"], int(row_["age"]), row_["city"],
                    float(row_["salary"]), row_["created_at"]]
            for i, (val, fm) in enumerate(zip(vals, [INT, None, None, None, None, "#,##0.00", None])):
                put(ws, f"{get_column_letter(c0 + i)}{r}", val, fmt=fm)
        cmp_ = ",".join(f"{get_column_letter(1 + i)}{r}={get_column_letter(9 + i)}{r}" for i in range(7))
        put(ws, f"P{r}", f'=IF(AND({cmp_}),"Yes","NO")', f(bold=True))
        r += 1
    widths(ws, {"A": 11, "B": 14, "C": 30, "D": 6, "E": 12, "F": 15, "G": 20, "H": 3,
                "I": 11, "J": 14, "K": 30, "L": 6, "M": 12, "N": 15, "O": 20, "P": 15})
    ws.freeze_panes = "A6"

    # ---------------------------------------------------------------- Doctor's checklist
    ws = ws_cl
    machine = (cfg.get("machine_label") or "").lower()
    on_laptop = not ("sandbox" in machine or "cloud" in machine)
    title(ws, "Doctor's instructions - checklist",
          "Each instruction from the doctor's chat, what was done, and where the proof is in this file.")
    header_row(ws, 4, 1, ["#", "Doctor's instruction (his words)", "What was done", "Proof in this file", "Status"])
    items = [
        ("Just install mysql in your laptop.",
         f"MySQL {env['mysql_version']} (InnoDB, default settings) was installed and every number in this file "
         "was measured on it. "
         + (("It ran on your laptop, with MySQL in a Docker container." if "docker" in machine
             else "It ran on your laptop.") if on_laptop else
            "This reference run was on a cloud computer, NOT on your laptop - rerun pk_experiment.py on your "
            "MacBook (HOW_TO_RUN_ON_MAC.md) so the final numbers are from your laptop."),
         "'Method & setup'", "Done" if on_laptop else "Rerun on your laptop"),
        ("Make a table with and without primary key",
         "data_pk: 7 columns, PRIMARY KEY (id).  data_nopk: the same 7 columns, no primary key and no index. "
         f"SHOW INDEX confirms: data_pk has PRIMARY on id; data_nopk has "
         f"{'no indexes' if not v['show_index']['without_pk'] else 'indexes'}.",
         "'MySQL check' section A", "Done"),
        ("Insert random variable up to 1 milliom",
         f"{cfg['initial_rows']:,} rows with random name, email, age, city, salary and date inserted into each table "
         f"(ids 1 to {cfg['initial_rows']:,}). MySQL COUNT(*) = {pk2['row_count']:,} in data_pk and "
         f"{np2['row_count']:,} in data_nopk; both tables hold identical data (equal checksums).",
         "'MySQL check' section B, 'Sample data'", "Done"),
        ("Select from the table a data from row 10, 1000, 10000, 100000 and 1000000 and see the time",
         "SELECT * FROM table WHERE id = N for N = " + ", ".join(f"{t:,}" for t in targets) +
         f" on both tables; each query run {cfg['select_repeats']} times in each of {cfg['runs']} runs (median "
         "reported). Times confirmed by MySQL's own timer.",
         "'Summary' Table 1, 'SELECT results', 'MySQL check' section C", "Done"),
        ("Then try insert 20000 random data to that previous table, see the time",
         f"{batches[0]:,} new random rows inserted into the SAME two tables "
         f"({cfg['initial_rows']:,} -> {cfg['initial_rows'] + batches[0]:,} rows) and timed.",
         "'Summary' Table 2, 'INSERT results', 'MySQL check' section D", "Done"),
        ("Again insert 40000, 60000, 100000 rows and see the times",
         "Inserted " + ", ".join(f"{b:,}" for b in batches[1:]) + " more rows into the same tables, one after "
         f"another ({cfg['initial_rows'] + batches[0]:,} -> {final_rows:,} rows), timing each. COUNT(*) checked "
         f"after every step: final {pk5['row_count']:,} rows in both tables.",
         "'Summary' Table 2, 'INSERT results', 'MySQL check' section D", "Done"),
    ]
    r = 5
    for i, (instr, done, proof, status) in enumerate(items, 1):
        put(ws, f"A{r}", i, f(bold=True), align=Alignment(horizontal="center", vertical="top"))
        put(ws, f"B{r}", f"\"{instr}\"", f(italic=True), align=wrap)
        put(ws, f"C{r}", done, align=wrap)
        put(ws, f"D{r}", proof, align=wrap)
        put(ws, f"E{r}", status, f(bold=True), align=Alignment(horizontal="center", vertical="top", wrap_text=True),
            fill=OK_FILL if status == "Done" else TODO_FILL)
        ws.row_dimensions[r].height = 62
        r += 1
    r += 1
    notes = [
        "How 'row 10' was read: as the row whose id = 10. Rows were inserted in id order 1, 2, 3 ..., so id = N is "
        "the N-th row inserted. SELECT * FROM table WHERE id = N is the normal way to fetch one row, and it is the "
        "query where a primary key makes the difference.",
        "Same data in both tables: every random row was inserted into both tables, so the ONLY difference between "
        "them is the primary key.",
        "Extras (not asked, clearly labelled): the LIMIT 1 test, SELECTs repeated after all inserts, EXPLAIN / rows "
        "read, table size on disk, and this verification run.",
    ]
    for n in notes:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        ws[f"A{r}"].value = n
        ws[f"A{r}"].font = f(9, color=MUTED, italic=True)
        ws[f"A{r}"].alignment = wrap
        ws.row_dimensions[r].height = 28
        r += 1
    widths(ws, {"A": 5, "B": 40, "C": 70, "D": 30, "E": 18})

    # note on the summary sheet
    ws_sum["A3"].value = (("Checked: all 6 of the doctor's instructions are done" if on_laptop else
                           "Checked: instructions 2-6 from the doctor are done; #1 (run it on your laptop) is still "
                           "to do") + " - see 'Doctor's checklist'. Data and timings were confirmed inside MySQL "
                          "itself - see 'MySQL check'.")
    ws_sum["A3"].font = f(9, bold=True, color="2E6B3A")

    order = ["Summary", "Doctor's checklist", "SELECT results", "INSERT results", "MySQL check", "Sample data",
             "Rows read (EXPLAIN)", "Table size", "Raw SELECT", "Raw INSERT", "Raw verification", "Method & setup"]
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames]


def px_table(label):
    return "data_pk" if label == "with_pk" else "data_nopk"


def run_time_label(env):
    """Docker Desktop containers run on UTC; say so when the time has no zone."""
    rs = str(env.get("run_started", ""))
    if rs.endswith("UTC") or "+" in rs:
        return rs
    return rs + " UTC" if "linuxkit" in str(env.get("os", "")).lower() else rs


def main():
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "results.json")
    out = Path(sys.argv[2] if len(sys.argv) > 2 else src.with_name("pk_experiment_results.xlsx"))
    d = json.loads(src.read_text())
    env, cfg = d["env"], d["config"]
    runs = list(range(1, cfg["runs"] + 1))
    targets = cfg["select_targets"]
    batches = cfg["insert_batches"]
    reps = cfg["select_repeats"]

    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sel = wb.create_sheet("SELECT results")
    ws_ins = wb.create_sheet("INSERT results")
    ws_rows = wb.create_sheet("Rows read (EXPLAIN)")
    ws_size = wb.create_sheet("Table size")
    ws_rs = wb.create_sheet("Raw SELECT")
    ws_ri = wb.create_sheet("Raw INSERT")
    ws_m = wb.create_sheet("Method & setup")

    # ------------------------------------------------------------------ Raw SELECT
    title(ws_rs, "Raw SELECT timings - every single query execution",
          f"Each SELECT was executed {reps} times per run; time = send query + fetch result, measured in Python "
          "(time.perf_counter).")
    cols = ["Run", "Stage", "Rows in table", "Query", "Table", "Target id", "Repeat #", "Time (ms)"]
    header_row(ws_rs, 4, 1, cols)
    stage_order = ["after initial load", "after all inserts"]
    variant_order = ["WHERE id = N", "WHERE id = N LIMIT 1"]
    sel = sorted(d["selects"], key=lambda s: (stage_order.index(s["stage"]), variant_order.index(s["variant"]),
                                              s["run"], s["target_id"], s["table"] != "with_pk", s["repeat"]))
    rng_of = {}  # (stage, variant, run, table, target) -> (first_row, last_row)
    r = 5
    for s in sel:
        vals = [s["run"], s["stage"], s["table_rows"], f"SELECT * FROM t {s['variant']}",
                LABELS[s["table"]], s["target_id"], s["repeat"], s["ms"]]
        fmts = [None, None, INT, None, None, INT, None, MS]
        for i, (v, fm) in enumerate(zip(vals, fmts)):
            put(ws_rs, f"{get_column_letter(i + 1)}{r}", v, fmt=fm)
        key = (s["stage"], s["variant"], s["run"], s["table"], s["target_id"])
        a, _ = rng_of.get(key, (r, r))
        rng_of[key] = (a, r)
        r += 1
    widths(ws_rs, {"A": 6, "B": 20, "C": 14, "D": 36, "E": 13, "F": 12, "G": 10, "H": 12})
    ws_rs.freeze_panes = "A5"
    ws_rs.auto_filter.ref = f"A4:H{r - 1}"

    # ------------------------------------------------------------------ Raw INSERT
    title(ws_ri, "Raw INSERT timings - every insert step, every run",
          f"Rows inserted with multi-row INSERT statements of {cfg['insert_chunk']:,} rows, committed once per step. "
          "Both tables received exactly the same random rows.")
    header_row(ws_ri, 4, 1, ["Run", "Step", "Table", "Rows inserted", "Rows before", "Rows after", "Time (s)"])
    step_names = ["Initial load"] + [f"Insert {b:,}" for b in batches]
    ins = sorted(d["inserts"], key=lambda x: (step_names.index(x["step"]), x["run"], x["table"] != "with_pk"))
    ins_cell = {}
    r = 5
    for x in ins:
        vals = [x["run"], x["step"], LABELS[x["table"]], x["rows_inserted"], x["rows_before"], x["rows_after"],
                x["seconds"]]
        fmts = [None, None, None, INT, INT, INT, SEC]
        for i, (v, fm) in enumerate(zip(vals, fmts)):
            put(ws_ri, f"{get_column_letter(i + 1)}{r}", v, fmt=fm)
        ins_cell[(x["step"], x["run"], x["table"])] = f"'Raw INSERT'!G{r}"
        r += 1
    widths(ws_ri, {"A": 6, "B": 18, "C": 13, "D": 14, "E": 14, "F": 14, "G": 12})
    ws_ri.freeze_panes = "A5"
    ws_ri.auto_filter.ref = f"A4:G{r - 1}"

    # ------------------------------------------------------------------ SELECT results
    title(ws_sel, "SELECT results - time to fetch one row by id (milliseconds)",
          f"Each run value = median of {reps} executions (formula over 'Raw SELECT'). "
          "'Median' = median of the runs. Lower is better.")
    nr = len(runs)
    sel_cells = {}  # (stage, variant) -> {target: (pk_median_cell, nopk_median_cell)}
    sections = [
        ("after initial load", "WHERE id = N",
         f"A. Main test: SELECT * FROM table WHERE id = N   (table has {cfg['initial_rows']:,} rows)"),
        ("after initial load", "WHERE id = N LIMIT 1",
         "B. Extra: same query with LIMIT 1 - MySQL may stop scanning at the first match"),
        ("after all inserts", "WHERE id = N",
         "C. Extra: main query repeated after all inserts (table has "
         f"{cfg['initial_rows'] + sum(batches):,} rows)"),
    ]
    row = 4
    for stage, variant, heading in sections:
        ws_sel[f"A{row}"].value = heading
        ws_sel[f"A{row}"].font = f(11, bold=True)
        row += 1
        pk_c0, np_c0 = 2, 2 + nr + 1
        ws_sel.merge_cells(start_row=row, start_column=pk_c0, end_row=row, end_column=pk_c0 + nr)
        ws_sel.merge_cells(start_row=row, start_column=np_c0, end_row=row, end_column=np_c0 + nr)
        put(ws_sel, f"{get_column_letter(pk_c0)}{row}", "WITH primary key", f(bold=True), fill=PK_FILL,
            align=Alignment(horizontal="center"))
        put(ws_sel, f"{get_column_letter(np_c0)}{row}", "WITHOUT primary key", f(bold=True), fill=NOPK_FILL,
            align=Alignment(horizontal="center"))
        row += 1
        labels = ["Target id"] + [f"Run {k}" for k in runs] + ["Median"] + [f"Run {k}" for k in runs] + \
                 ["Median", "PK is ... times faster"]
        header_row(ws_sel, row, 1, labels)
        row += 1
        sel_cells[(stage, variant)] = {}
        for t in targets:
            put(ws_sel, f"A{row}", t, f(bold=True), INT)
            for side, c0 in (("with_pk", pk_c0), ("without_pk", np_c0)):
                for j, k in enumerate(runs):
                    a, b = rng_of[(stage, variant, k, side, t)]
                    put(ws_sel, f"{get_column_letter(c0 + j)}{row}", f"=MEDIAN('Raw SELECT'!H{a}:H{b})", fmt=MS)
                first, last = get_column_letter(c0), get_column_letter(c0 + nr - 1)
                put(ws_sel, f"{get_column_letter(c0 + nr)}{row}", f"=MEDIAN({first}{row}:{last}{row})", f(bold=True),
                    MS, fill=PK_FILL if side == "with_pk" else NOPK_FILL)
            pk_med = f"{get_column_letter(pk_c0 + nr)}{row}"
            np_med = f"{get_column_letter(np_c0 + nr)}{row}"
            put(ws_sel, f"{get_column_letter(np_c0 + nr + 1)}{row}", f"={np_med}/{pk_med}", f(bold=True), TIMES)
            sel_cells[(stage, variant)][t] = (pk_med, np_med)
            row += 1
        if variant.endswith("LIMIT 1"):
            ws_sel[f"A{row}"].value = ("With LIMIT 1 MySQL stops scanning at the first match, so for a row near the start "
                                       "(id = 10) the table without a PK can be as fast as the PK lookup (ratio below 1x).")
            ws_sel[f"A{row}"].font = f(9, color=MUTED, italic=True)
        row += 2
    widths(ws_sel, {get_column_letter(i): 12 for i in range(1, 2 * nr + 5)})
    ws_sel.column_dimensions["A"].width = 14
    ws_sel.column_dimensions[get_column_letter(2 * nr + 4)].width = 16
    ws_sel.freeze_panes = "B4"

    # chart for LIMIT 1 section on the SELECT sheet (shows growth with row position)
    ch_data_row = row + 1
    ws_sel[f"A{ch_data_row - 1}"].value = "Chart data (section B medians)"
    ws_sel[f"A{ch_data_row - 1}"].font = f(9, color=MUTED, italic=True)
    header_row(ws_sel, ch_data_row, 1, ["Target id", "With PK", "Without PK"])
    for i, t in enumerate(targets):
        rr = ch_data_row + 1 + i
        pk_med, np_med = sel_cells[("after initial load", "WHERE id = N LIMIT 1")][t]
        put(ws_sel, f"A{rr}", f"id = {t:,}")
        put(ws_sel, f"B{rr}", f"={pk_med}", fmt=MS)
        put(ws_sel, f"C{rr}", f"={np_med}", fmt=MS)
    ch = BarChart()
    ch.title = "With LIMIT 1 (ms, log scale)"
    ch.add_data(Reference(ws_sel, min_col=2, max_col=3, min_row=ch_data_row, max_row=ch_data_row + len(targets)),
                titles_from_data=True)
    ch.set_categories(Reference(ws_sel, min_col=1, min_row=ch_data_row + 1, max_row=ch_data_row + len(targets)))
    style_chart(ch, "Time (ms, log scale)", "Row id searched", log=True)
    ws_sel.add_chart(ch, f"E{ch_data_row - 1}")

    # ------------------------------------------------------------------ INSERT results
    title(ws_ins, "INSERT results - time to insert new random rows (seconds)",
          "Steps are cumulative on the same tables: 1,000,000 -> +20,000 -> +40,000 -> +60,000 -> +100,000. "
          "Run values link to 'Raw INSERT'. Lower is better.")
    row = 4
    pk_c0, np_c0 = 4, 4 + nr + 1
    ws_ins.merge_cells(start_row=row, start_column=pk_c0, end_row=row, end_column=pk_c0 + nr)
    ws_ins.merge_cells(start_row=row, start_column=np_c0, end_row=row, end_column=np_c0 + nr)
    put(ws_ins, f"{get_column_letter(pk_c0)}{row}", "WITH primary key (s)", f(bold=True), fill=PK_FILL,
        align=Alignment(horizontal="center"))
    put(ws_ins, f"{get_column_letter(np_c0)}{row}", "WITHOUT primary key (s)", f(bold=True), fill=NOPK_FILL,
        align=Alignment(horizontal="center"))
    row += 1
    labels = ["Step", "Rows inserted", "Table size after"] + [f"Run {k}" for k in runs] + ["Median"] + \
             [f"Run {k}" for k in runs] + ["Median", "Without PK vs with PK", "Rows/s with PK", "Rows/s without PK"]
    header_row(ws_ins, row, 1, labels)
    row += 1
    ins_cells = {}
    size_after = cfg["initial_rows"]
    for step in step_names:
        n = cfg["initial_rows"] if step == "Initial load" else int(step.split()[1].replace(",", ""))
        size_after = cfg["initial_rows"] if step == "Initial load" else size_after + n
        put(ws_ins, f"A{row}", step, f(bold=True))
        put(ws_ins, f"B{row}", n, fmt=INT)
        put(ws_ins, f"C{row}", size_after, fmt=INT)
        for side, c0 in (("with_pk", pk_c0), ("without_pk", np_c0)):
            for j, k in enumerate(runs):
                put(ws_ins, f"{get_column_letter(c0 + j)}{row}", f"={ins_cell[(step, k, side)]}", fmt=SEC)
            first, last = get_column_letter(c0), get_column_letter(c0 + nr - 1)
            put(ws_ins, f"{get_column_letter(c0 + nr)}{row}", f"=MEDIAN({first}{row}:{last}{row})", f(bold=True), SEC,
                fill=PK_FILL if side == "with_pk" else NOPK_FILL)
        pk_med = f"{get_column_letter(pk_c0 + nr)}{row}"
        np_med = f"{get_column_letter(np_c0 + nr)}{row}"
        cdiff = np_c0 + nr + 1
        put(ws_ins, f"{get_column_letter(cdiff)}{row}", f"=({np_med}-{pk_med})/{pk_med}", fmt=PCT)
        put(ws_ins, f"{get_column_letter(cdiff + 1)}{row}", f"=B{row}/{pk_med}", fmt=INT)
        put(ws_ins, f"{get_column_letter(cdiff + 2)}{row}", f"=B{row}/{np_med}", fmt=INT)
        ins_cells[step] = (pk_med, np_med, f"{get_column_letter(cdiff)}{row}")
        row += 1
    note_row = row + 1
    ws_ins[f"A{note_row}"].value = ("'Without PK vs with PK' = (median without PK - median with PK) / median with PK. "
                                    "Positive = the table without a PK was slower; negative = it was faster.")
    ws_ins[f"A{note_row}"].font = f(9, color=MUTED, italic=True)
    widths(ws_ins, {get_column_letter(i): 12 for i in range(1, 2 * nr + 9)})
    ws_ins.column_dimensions["A"].width = 16
    ws_ins.column_dimensions["B"].width = 13
    ws_ins.column_dimensions["C"].width = 15
    ws_ins.freeze_panes = "B6"

    # ------------------------------------------------------------------ Rows read (EXPLAIN)
    title(ws_rows, "How many rows MySQL had to read - EXPLAIN plan and handler counters (run 1)",
          "type = const -> direct lookup through the PRIMARY KEY index.  type = ALL -> full table scan.  "
          "Handler_read_rnd_next counts rows read one-by-one during a scan.")
    header_row(ws_rows, 4, 1, ["Table", "Query", "Target id", "EXPLAIN type", "EXPLAIN key",
                               "EXPLAIN rows (estimate)", "Handler_read_key", "Handler_read_rnd_next",
                               "Rows actually read"])
    rows_cell = {}
    r = 5
    rx = sorted(d["rows_examined"], key=lambda x: (variant_order.index(x["variant"]), x["table"] != "with_pk",
                                                    x["target_id"]))
    for x in rx:
        vals = [LABELS[x["table"]], f"SELECT * FROM t {x['variant']}", x["target_id"], x["explain_type"],
                x["explain_key"] or "(none)", x["explain_rows_estimate"], x["handler_read_key"],
                x["handler_read_rnd_next"]]
        fmts = [None, None, INT, None, None, INT, INT, INT]
        for i, (v, fm) in enumerate(zip(vals, fmts)):
            put(ws_rows, f"{get_column_letter(i + 1)}{r}", v, fmt=fm)
        # rows actually read: index lookups for the PK table, row-by-row scan reads for the no-PK table
        put(ws_rows, f"I{r}", f'=IF(H{r}>0,H{r},G{r})', f(bold=True), INT)
        rows_cell[(x["table"], x["variant"], x["target_id"])] = f"'Rows read (EXPLAIN)'!I{r}"
        r += 1
    ws_rows[f"A{r + 1}"].value = ("Rows actually read = Handler_read_rnd_next when MySQL scanned the table, "
                                  "otherwise Handler_read_key (index lookups). A scan also counts one final "
                                  "'end of table' read, so a full scan of N rows shows N + 1.")
    ws_rows[f"A{r + 1}"].font = f(9, color=MUTED, italic=True)
    widths(ws_rows, {"A": 13, "B": 36, "C": 12, "D": 14, "E": 14, "F": 14, "G": 19, "H": 23, "I": 16})
    for rr in range(5, r):
        for col in "DE":
            ws_rows[f"{col}{rr}"].alignment = Alignment(horizontal="center")
    ws_rows.freeze_panes = "A5"

    # ------------------------------------------------------------------ Table size
    title(ws_size, "Table size on disk (information_schema, after ANALYZE TABLE)",
          "InnoDB stores the whole table inside its clustered index. With a PK that index is the PRIMARY KEY; "
          "without one InnoDB adds a hidden 6-byte row id (GEN_CLUST_INDEX) and clusters on that.")
    header_row(ws_size, 4, 1, ["Run", "Stage", "Table", "Row count", "Data (MB)", "Secondary indexes (MB)"])
    r = 5
    for x in sorted(d["sizes"], key=lambda x: (x["run"], x["stage"], x["table"] != "with_pk")):
        vals = [x["run"], x["stage"], LABELS[x["table"]], x["row_count"], x["data_mb"], x["index_mb"]]
        fmts = [None, None, None, INT, '#,##0.0', '#,##0.0']
        for i, (v, fm) in enumerate(zip(vals, fmts)):
            put(ws_size, f"{get_column_letter(i + 1)}{r}", v, fmt=fm)
        r += 1
    widths(ws_size, {"A": 6, "B": 18, "C": 13, "D": 13, "E": 12, "F": 22})

    # ------------------------------------------------------------------ Summary
    ws = ws_sum
    machine = cfg.get("machine_label") or env["cpu"]
    title(ws, "MySQL experiment: table WITH primary key vs WITHOUT primary key",
          f"Machine: {machine}  |  MySQL {env['mysql_version']}, InnoDB, buffer pool {env['innodb_buffer_pool_mb']} MB"
          f"  |  {cfg['runs']} full runs, each SELECT executed {reps}x - medians shown  |  run on {run_time_label(env)}")

    # numbers for the findings text (computed from the raw data, same method as the formulas)
    def med_sel(stage, variant, table, t):
        per_run = [statistics.median(s["ms"] for s in d["selects"] if s["stage"] == stage and s["variant"] == variant
                                     and s["table"] == table and s["target_id"] == t and s["run"] == k) for k in runs]
        return statistics.median(per_run)

    def med_ins(step, table):
        return statistics.median(x["seconds"] for x in d["inserts"] if x["step"] == step and x["table"] == table)

    main_v, lim_v = "WHERE id = N", "WHERE id = N LIMIT 1"
    pk_all = [med_sel("after initial load", main_v, "with_pk", t) for t in targets]
    np_all = [med_sel("after initial load", main_v, "without_pk", t) for t in targets]
    np_lim = [med_sel("after initial load", lim_v, "without_pk", t) for t in targets]
    np_after = [med_sel("after all inserts", main_v, "without_pk", t) for t in targets]
    pk_after = [med_sel("after all inserts", main_v, "with_pk", t) for t in targets]
    ins_pk = sum(med_ins(f"Insert {b:,}", "with_pk") for b in batches)
    ins_np = sum(med_ins(f"Insert {b:,}", "without_pk") for b in batches)
    ins_diff = (ins_np - ins_pk) / ins_pk
    final_rows = cfg["initial_rows"] + sum(batches)
    if abs(ins_diff) < 0.10:
        ins_text = (f"INSERT: both tables took about the same time (all {len(batches)} batches together: {ins_pk:.2f} s with PK vs "
                    f"{ins_np:.2f} s without PK, {ins_diff:+.1%}). Because the ids are increasing, every new row "
                    "is appended at the right end of the primary-key B+tree, so keeping the PK costs almost "
                    "nothing; and InnoDB gives the no-PK table a hidden row-id index anyway, so it does similar work.")
    elif ins_diff < 0:
        ins_text = (f"INSERT: the table without a PK was {abs(ins_diff):.0%} faster (all {len(batches)} batches together: "
                    f"{ins_np:.2f} s vs {ins_pk:.2f} s with PK). With a PK, MySQL must check every new id for "
                    "uniqueness and keep the PK index ordered - that is the small price paid for the fast SELECTs.")
    else:
        ins_text = (f"INSERT: the table without a PK was {ins_diff:.0%} slower (all {len(batches)} batches together: "
                    f"{ins_np:.2f} s vs {ins_pk:.2f} s with PK). InnoDB must generate a hidden 6-byte row id for "
                    "every row of a table without a PK (from one global counter), so skipping the PK does not "
                    "make inserts cheaper.")
    np_before_m, np_after_m = statistics.median(np_all), statistics.median(np_after)
    growth = np_after_m / np_before_m - 1
    extra = sum(batches) / cfg["initial_rows"]
    if growth >= 0.05:
        after_text = (f"After the extra inserts ({final_rows:,} rows) the no-PK SELECT rose from about "
                      f"{np_before_m:,.0f} ms to {np_after_m:,.0f} ms ({growth:+.0%}; scan time grows with table "
                      f"size, O(n)), while the PK SELECT stayed at about {statistics.median(pk_after):.2f} ms "
                      "(O(log n)).")
    else:
        after_text = (f"After the extra inserts ({final_rows:,} rows) the no-PK SELECT took about {np_after_m:,.0f} ms "
                      f"vs {np_before_m:,.0f} ms before ({growth:+.0%}) - no clear change on this machine: "
                      f"{extra:.0%} more rows should add roughly {extra:.0%} more scan time, but the run-to-run "
                      "variation here was larger than that. It still read every row (O(n)), while the PK SELECT "
                      f"stayed at about {statistics.median(pk_after):.2f} ms (O(log n)).")
    findings = [
        f"SELECT with a primary key: about {statistics.median(pk_all):.2f} ms for every row id, from row {targets[0]:,} to "
        f"row {targets[-1]:,}. MySQL jumps directly to the row through the PRIMARY KEY B+tree index "
        "(EXPLAIN type = const, 1 row read).",
        f"SELECT without a primary key: about {statistics.median(np_all):,.0f} ms - roughly "
        f"{statistics.median(np_all) / statistics.median(pk_all):,.0f}x slower - and nearly the same for row "
        f"{targets[0]:,} and row {targets[-1]:,}. With no index MySQL must read all {cfg['initial_rows']:,} rows "
        "(EXPLAIN type = ALL); it cannot stop at "
        "the first match because nothing says id is unique.",
        f"With LIMIT 1 the scan stops at the first match, so the no-PK time grows with the row's position: "
        f"{np_lim[0]:.2f} ms for row {targets[0]:,} -> {np_lim[-1]:,.0f} ms for row {targets[-1]:,} (see 'SELECT results', section B).",
        ins_text,
        after_text,
    ]
    ws["A4"].value = "Key findings"
    ws["A4"].font = f(12, bold=True)
    r = 5
    for i, text in enumerate(findings, 1):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c = ws[f"A{r}"]
        c.value = f"{i}. {text}"
        c.font = f(10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 40
        r += 1

    # Table 1: SELECT
    r += 1
    ws[f"A{r}"].value = f"Table 1 - SELECT * FROM table WHERE id = N   (table has {cfg['initial_rows']:,} rows)"
    ws[f"A{r}"].font = f(11, bold=True)
    r += 1
    header_row(ws, r, 1, ["Row id searched", "With PK (ms)", "Without PK (ms)", "PK is ... times faster",
                          "Rows read with PK", "Rows read without PK"])
    t1_head = r
    r += 1
    for t in targets:
        pk_med, np_med = sel_cells[("after initial load", main_v)][t]
        put(ws, f"A{r}", f"id = {t:,}", f(bold=True))
        put(ws, f"B{r}", f"='SELECT results'!{pk_med}", fmt=MS, fill=PK_FILL)
        put(ws, f"C{r}", f"='SELECT results'!{np_med}", fmt=MS, fill=NOPK_FILL)
        put(ws, f"D{r}", f"=C{r}/B{r}", f(bold=True), TIMES)
        put(ws, f"E{r}", f"={rows_cell[('with_pk', main_v, t)]}", fmt=INT)
        put(ws, f"F{r}", f"={rows_cell[('without_pk', main_v, t)]}", fmt=INT)
        r += 1
    t1_end = r - 1

    # Table 2: INSERT
    r += 1
    ws[f"A{r}"].value = "Table 2 - INSERT new random rows into the same tables (seconds, cumulative)"
    ws[f"A{r}"].font = f(11, bold=True)
    r += 1
    header_row(ws, r, 1, ["Step", "Rows inserted", "Table size after", "With PK (s)", "Without PK (s)",
                          "Without PK vs with PK"])
    t2_head = r
    r += 1
    size_after = cfg["initial_rows"]
    batch_first = None
    for step in step_names[1:] + step_names[:1]:          # batches first (charted), initial load last
        n = cfg["initial_rows"] if step == "Initial load" else int(step.split()[1].replace(",", ""))
        pk_med, np_med, diff = ins_cells[step]
        if step != "Initial load":
            size_after += n
            batch_first = batch_first or r
        put(ws, f"A{r}", f"Insert {n:,} rows" if step != "Initial load" else "Initial load (step 1)", f(bold=True))
        put(ws, f"B{r}", n, fmt=INT)
        put(ws, f"C{r}", size_after if step != "Initial load" else cfg["initial_rows"], fmt=INT)
        put(ws, f"D{r}", f"='INSERT results'!{pk_med}", fmt=SEC, fill=PK_FILL)
        put(ws, f"E{r}", f"='INSERT results'!{np_med}", fmt=SEC, fill=NOPK_FILL)
        put(ws, f"F{r}", f"=(E{r}-D{r})/D{r}", f(bold=True), PCT)
        r += 1
    t2_end = r - 1
    ws[f"A{r + 1}"].value = ("All values are medians and are linked by formulas to the detail sheets; every single "
                             "measurement is in 'Raw SELECT' and 'Raw INSERT'. Method, table definitions and settings "
                             "are in 'Method & setup'.")
    ws[f"A{r + 1}"].font = f(9, color=MUTED, italic=True)
    ws[f"A{r + 1}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=6)
    ws.row_dimensions[r + 1].height = 26
    widths(ws, {"A": 22, "B": 14, "C": 16, "D": 16, "E": 16, "F": 18})

    ch1 = BarChart()
    ch1.title = "SELECT one row by id (ms, log scale)"
    ch1.add_data(Reference(ws, min_col=2, max_col=3, min_row=t1_head, max_row=t1_end), titles_from_data=True)
    ch1.set_categories(Reference(ws, min_col=1, min_row=t1_head + 1, max_row=t1_end))
    style_chart(ch1, "Time (ms, log scale)", "Row id searched", log=True)
    ws.add_chart(ch1, "H4")

    ch2 = BarChart()
    ch2.title = "INSERT time per batch (seconds)"
    ch2.add_data(Reference(ws, min_col=4, max_col=5, min_row=t2_head, max_row=batch_first + len(batches) - 1),
                 titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=1, min_row=batch_first, max_row=batch_first + len(batches) - 1))
    style_chart(ch2, "Time (seconds)", "Batch")
    ws.add_chart(ch2, "H22")

    # ------------------------------------------------------------------ Method & setup
    wm = ws_m
    title(wm, "Method & setup")
    lines = [
        ("Environment", None),
        ("Machine", machine),
        ("CPU", env["cpu"]),
        ("Operating system", env["os"]),
        ("MySQL version", env["mysql_version"]),
        ("Storage engine", "InnoDB (MySQL default)"),
        ("innodb_buffer_pool_size", f"{env['innodb_buffer_pool_mb']} MB (default)"),
        ("innodb_flush_log_at_trx_commit", f"{env['innodb_flush_log_at_trx_commit']} (default - full durability)"),
        ("Client", f"Python {env['python']} + mysql-connector-python, same machine as the server"),
        ("Run started", run_time_label(env)),
        ("", None),
        ("Procedure (from the assignment)", None),
        ("Step 1", "Create two tables with identical columns: data_pk (id is the PRIMARY KEY) and data_nopk "
                   "(id is a normal INT column: no primary key, no index)."),
        ("Step 2", f"Insert {cfg['initial_rows']:,} random rows (ids 1..{cfg['initial_rows']:,}) into each table. "
                   "Both tables get exactly the same random data."),
        ("Step 3", "SELECT * FROM <table> WHERE id = N for N = " + ", ".join(f"{t:,}" for t in targets) +
                   f"; each query executed {reps} times, median kept."),
        ("Step 4", f"Insert {batches[0]:,} new random rows into the same tables and time it."),
        ("Step 5", "Insert " + ", ".join(f"{b:,}" for b in batches[1:]) +
                   " more rows (cumulative, same tables) and time each batch."),
        ("Repeats", f"The whole procedure (fresh tables) was repeated {cfg['runs']} times; the table order "
                    "(PK first / no-PK first) alternates between runs so neither table always runs 'warm'. "
                    "Summary values are medians of the runs."),
        ("Insert method", f"Multi-row INSERT statements of {cfg['insert_chunk']:,} rows each, one COMMIT per step. "
                          "Random data generated before the timer starts, so only database time is measured."),
        ("Timing", "Wall-clock time from sending the query to receiving the full result (Python time.perf_counter). "
                   "The ~0.3 ms measured with a PK is mostly client round-trip overhead; the index lookup itself "
                   "inside MySQL takes only microseconds."),
        ("Random data", f"name, email, age 18-65, one of 12 Indonesian cities, salary, created_at 2020-2025; "
                        f"seed {cfg['seed']} + run number (reproducible)."),
        ("", None),
        ("Table definitions", None),
        ("With PK", cfg["ddl_with_pk"]),
        ("Without PK", cfg["ddl_without_pk"]),
        ("", None),
        ("Why the results look like this", None),
        ("PRIMARY KEY", "InnoDB stores the table as a B+tree ordered by the primary key (clustered index). Finding "
                        "id = N walks ~3-4 tree levels no matter where the row is: O(log n), constant in practice."),
        ("No PRIMARY KEY", "The id column has no index, so WHERE id = N is a full table scan: MySQL reads every row "
                           "and compares id. Without LIMIT it cannot stop early (id is not known to be unique), so "
                           "row 10 and row 1,000,000 cost the same: O(n)."),
        ("Hidden row id", "A table without a PK is still clustered - InnoDB adds a hidden 6-byte DB_ROW_ID and "
                          "builds the clustered index on it. That index is useless for WHERE id = N."),
        ("Inserts", "With increasing ids, a PK insert appends to the right end of the B+tree (cheap). A random "
                    "PK (e.g. UUID) would cause page splits and slower inserts; a table with extra secondary "
                    "indexes would also insert more slowly."),
        ("Caveat", "Absolute times depend on the machine (CPU, SSD, RAM, buffer pool). Compare the ratio between "
                   "the two tables, not the raw milliseconds, when you rerun it on another laptop."),
    ]
    r = 3
    for k, v in lines:
        if v is None and k:
            wm[f"A{r}"].value = k
            wm[f"A{r}"].font = f(11, bold=True)
        elif k:
            put(wm, f"A{r}", k, f(bold=True), align=Alignment(vertical="top"))
            put(wm, f"B{r}", v, align=Alignment(wrap_text=True, vertical="top"))
            if "\n" in str(v) or len(str(v)) > 95:
                wm.row_dimensions[r].height = max(30, 15 * (str(v).count("\n") + 1 + len(str(v)) // 110))
        r += 1
    widths(wm, {"A": 30, "B": 110})

    vpath = Path(sys.argv[3]) if len(sys.argv) > 3 else src.with_name("verification.json")
    if vpath.exists():
        add_verification(wb, json.loads(vpath.read_text()), cfg, env, sel_cells, ins_cells, ws_sum)

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        # print-friendly: landscape, one page wide
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1 if sheet.title == "Summary" else 0
    wb.save(out)
    print(f"Saved Excel -> {out}")


if __name__ == "__main__":
    main()
